"""Exercise real HTTP, a separate worker and Excel using an isolated temporary DB.

Run: python qa/smoke_service.py [--real-data]
No production users, database, source files or services are changed.
"""
from __future__ import annotations

import argparse
import copy
import http.cookiejar
import io
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"


def stop_process(process):
    """Stop only the tree started by this smoke test, including Windows launchers."""
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       creationflags=subprocess.CREATE_NO_WINDOW, check=False)
    else:
        process.terminate()
    try:
        process.wait(timeout=15)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=5)


class Client:
    def __init__(self, origin):
        self.origin = origin
        self.csrf = ""
        self.cookies = http.cookiejar.CookieJar()
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(self.cookies))

    def request(self, path, method="GET", data=None, expected=200, csrf=True, headers=None):
        outgoing = {"Origin": self.origin, **(headers or {})}
        if data is not None:
            outgoing["Content-Type"] = "application/json"
        if csrf and self.csrf:
            outgoing["X-CSRF-Token"] = self.csrf
        request = urllib.request.Request(
            self.origin + "/api" + path,
            data=json.dumps(data).encode() if data is not None else None,
            method=method, headers=outgoing,
        )
        try:
            response = self.opener.open(request, timeout=180)
        except urllib.error.HTTPError as error:
            response = error
        content = response.read()
        assert response.status == expected, (method, path, response.status, content[:500])
        if "json" in response.headers.get("Content-Type", ""):
            return json.loads(content)
        return content, response.headers

    def login(self, username, password):
        value = self.request("/auth/login", "POST", {"username": username, "password": password})
        self.csrf = value["csrf_token"]
        return value["user"]


def await_job(client, job_id):
    deadline = time.monotonic() + 620
    while time.monotonic() < deadline:
        job = client.request("/jobs/" + job_id)
        if job["status"] == "completed":
            return client.request("/orders/" + job["order_id"])
        assert job["status"] not in {"failed", "cancelled"}, job
        time.sleep(0.5)
    raise AssertionError("Calculation did not complete within its time limit")


def verify_excel(client, order, approved_only):
    from openpyxl import load_workbook

    binary, headers = client.request(
        "/orders/" + order["id"] + "/export", "POST",
        {"revision": order["revision"], "approved_only": approved_only},
    )
    assert headers.get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert binary.startswith(b"PK\x03\x04")
    workbook = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    assert workbook.sheetnames == ["Заказ поставщикам", "Параметры расчёта"]
    rows = list(workbook.worksheets[0].values)
    exported = [dict(zip(rows[0], row)) for row in rows[1:]]
    decisions = {row["line_id"]: row for row in order["decisions"]}
    selected = [(line, decisions[line["line_id"]])
                for group in order["calculation"]["groups"] for line in group["lines"]
                if decisions[line["line_id"]]["quantity"] > 0
                and (not approved_only or decisions[line["line_id"]]["approved"])]
    assert len(exported) == len(selected) > 0
    for actual, (line, decision) in zip(exported, selected):
        assert str(actual["Артикул"]) == line["sku"]
        assert (actual["Артикул поставщика"] or "") == (line.get("supplier_sku") or "")
        assert actual["К заказу, ед"] == decision["quantity"]
        assert actual["Рекомендовано, ед"] == line["recommended_qty"]
        assert actual["Склад"] == (line["warehouse"] or order["calculation"].get("warehouse") or "Все")
        assert actual["Ед. изм."] == line["unit"]
        assert actual["Утверждено"] == ("Да" if decision["approved"] else "Нет")
    parameters = dict(list(workbook.worksheets[1].values)[1:])
    assert parameters["ID расчёта"] == order["calculation"]["calculation_id"]
    assert parameters["Количество позиций"] == len(selected)
    workbook.close()
    return len(selected)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--real-data", action="store_true", help="Use the repository's actual read-only Excel sources")
    args = parser.parse_args()
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    origin = f"http://127.0.0.1:{port}"
    with tempfile.TemporaryDirectory(prefix="umytpa-http-qa-") as temporary:
        directory = Path(temporary)
        os.environ.update({
            "APP_ENV": "development", "APP_ORIGIN": origin,
            "APP_DB_PATH": str(directory / "orders.sqlite3"),
            "ORDER_DB_PATH": str(directory / "orders.sqlite3"),
            "DATA_SOURCE": "excel" if args.real_data else "synthetic",
            "DATA_DIR": str(ROOT), "OPENAI_API_KEY": "", "CORS_ORIGINS": "",
            "PYTHONUTF8": "1", "PYTHONUNBUFFERED": "1",
        })
        sys.path.insert(0, str(BACKEND))
        from app.storage import migrate
        from app.repositories import create_user
        migrate()
        password = secrets.token_urlsafe(24)
        create_user("qa_admin", password, "admin")
        children, logs = [], []

        def launch(module, *arguments):
            log = open(directory / (module.replace(".", "_") + f"-{len(logs)}.log"), "w", encoding="utf-8")
            logs.append(log)
            process = subprocess.Popen(
                [sys.executable, "-m", module, *arguments], cwd=BACKEND,
                stdout=log, stderr=subprocess.STDOUT,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
            )
            children.append(process)
            return process

        def wait_ready():
            deadline = time.monotonic() + 45
            while time.monotonic() < deadline:
                assert api.poll() is None, "The newly launched API exited before readiness"
                assert worker.poll() is None, "The worker exited before readiness"
                try:
                    with urllib.request.urlopen(origin + "/api/ready", timeout=2) as response:
                        if response.status == 200:
                            assert api.poll() is None, "Readiness came from a different API process"
                            return
                except (OSError, urllib.error.URLError):
                    pass
                time.sleep(0.3)
            raise AssertionError("API and worker did not become ready")

        try:
            api = launch("uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port), "--no-access-log")
            worker = launch("app.worker")
            wait_ready()
            anonymous, admin, manager, colleague = [Client(origin) for _ in range(4)]
            anonymous.request("/orders", expected=401)
            admin.login("qa_admin", password)
            manager_user = admin.request("/users", "POST", {"username": "qa_manager", "password": password, "role": "manager"}, expected=201)
            admin.request("/users", "POST", {"username": "qa_colleague", "password": password, "role": "manager"}, expected=201)
            manager.login("qa_manager", password)
            colleague.login("qa_colleague", password)
            manager.request("/users", expected=403)
            admin.request("/users", "POST", {"username": "not_created", "password": password, "role": "manager"}, csrf=False, expected=403)
            admin.request("/users", "POST", {"username": "not_created", "password": password, "role": "manager"}, headers={"Origin": "https://untrusted.invalid"}, expected=403)
            metadata = manager.request("/meta")
            assert metadata["sku_count"] > 0 and not metadata["capabilities"]["llm_available"]
            assert "defaults" in metadata
            print("PASS: authentication, roles, CSRF, same-origin checks and metadata", flush=True)

            key = secrets.token_urlsafe(24)
            params = {"explain": False}
            headers = {"Prefer": "respond-async", "Idempotency-Key": key}
            job = manager.request("/recommend", "POST", params, expected=202, headers=headers)
            repeated = manager.request("/recommend", "POST", params, expected=202, headers=headers)
            assert repeated["job_id"] == job["job_id"]
            manager.request("/recommend", "POST", {"explain": False, "review_period_days": 7}, expected=409, headers=headers)
            colleague.request("/jobs/" + job["job_id"], expected=403)
            order = await_job(manager, job["job_id"])
            order_path = "/orders/" + order["id"]
            colleague.request(order_path, expected=403)
            colleague.request(order_path + "/export", "POST", {"revision": order["revision"], "approved_only": False}, expected=403)
            assert colleague.request("/orders")["total"] == 0
            assert admin.request("/orders")["total"] == 1
            assert manager.request("/orders")["items"][0]["id"] == order["id"]
            lines = [line for group in order["calculation"]["groups"] for line in group["lines"]]
            assert len(lines) > 0
            line = next(row for row in lines if row["recommended_qty"] > 0)
            index = next(i for i, row in enumerate(order["decisions"]) if row["line_id"] == line["line_id"])
            changes = copy.deepcopy(order["decisions"])
            changes[index]["approved"] = True
            order = manager.request(order_path, "PUT", {"revision": order["revision"], "lines": changes})
            old_revision = order["revision"]
            changes[index]["quantity"] += line["pack_size"] or 1
            order = manager.request(order_path, "PUT", {"revision": old_revision, "lines": changes})
            current = next(row for row in order["decisions"] if row["line_id"] == line["line_id"])
            assert current["approved"] is False
            manager.request(order_path, "PUT", {"revision": old_revision, "lines": changes}, expected=409)
            manager.request(order_path + "/export", "POST", {"revision": old_revision, "approved_only": False}, expected=409)
            invalid = copy.deepcopy(order["decisions"])
            invalid[0]["quantity"] = -1
            manager.request(order_path, "PUT", {"revision": order["revision"], "lines": invalid}, expected=422)
            changes = copy.deepcopy(order["decisions"])
            next(row for row in changes if row["line_id"] == line["line_id"])["approved"] = True
            order = manager.request(order_path, "PUT", {"revision": order["revision"], "lines": changes})
            approved_count = verify_excel(manager, order, True)
            draft_count = verify_excel(manager, order, False)
            assert approved_count == 1
            events = manager.request(order_path + "/events")["items"]
            assert any(event["action"] == "exported" for event in events)
            assert password not in json.dumps(events)
            print(f"PASS: worker calculation, deduplication, ownership, revisions, approvals; Excel {draft_count} draft / {approved_count} approved rows", flush=True)

            stop_process(api)
            api = launch("uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port), "--no-access-log")
            wait_ready()
            restored = manager.request(order_path)
            assert restored["revision"] == order["revision"] and restored["decisions"] == order["decisions"]
            cancel_job = colleague.request("/recommend", "POST", params, expected=202, headers={"Prefer": "respond-async", "Idempotency-Key": secrets.token_urlsafe(24)})
            cancelled = colleague.request("/jobs/" + cancel_job["job_id"] + "/cancel", "POST", {})
            assert cancelled["status"] == "cancelled"
            time.sleep(2)
            assert colleague.request("/orders")["total"] == 0
            admin.request("/users/" + manager_user["id"], "PATCH", {"active": False})
            manager.request(order_path, expected=401)
            colleague.request("/auth/logout", "POST", {}, expected=204)
            colleague.request("/auth/session", expected=401)
            for _ in range(5):
                anonymous.request("/auth/login", "POST", {"username": "unknown_user", "password": password}, expected=401)
            anonymous.request("/auth/login", "POST", {"username": "unknown_user", "password": password}, expected=429)
            print("PASS: process restart persistence, cancellation, blocked sessions, logout and login throttling", flush=True)
            print("SUCCESS: real HTTP integration with " + ("supplied Excel sources" if args.real_data else "synthetic fixtures"), flush=True)
        except Exception:
            for log in logs:
                log.flush()
                print(Path(log.name).read_text(encoding="utf-8", errors="replace")[-6000:], file=sys.stderr)
            raise
        finally:
            for process in reversed(children):
                stop_process(process)
            for log in logs:
                log.close()


if __name__ == "__main__":
    main()
