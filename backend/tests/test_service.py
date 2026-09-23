"""Security, concurrent order changes, durable jobs and migration guarantees."""
from __future__ import annotations

import io
import json
import os
import tempfile
import time
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import jobs
from app.core.explain import build_explanation, explanation_budget
from app.core.snapshots import save_snapshot
from app.main import app
from app.repositories import create_order, create_user, get_order, public_metadata
from app.security import COOKIE_NAME
from app.storage import connection, migrate
from app.worker import heartbeat, register
from tests.test_export import example_response


class ServiceTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        path = self.directory.name + "/service.sqlite3"
        env = patch.dict(os.environ, {"APP_DB_PATH": path, "ORDER_DB_PATH": path,
                                     "APP_ENV": "development", "APP_ORIGIN": "", "JOB_TIMEOUT_SECONDS": "600"})
        env.start()
        self.addCleanup(env.stop)
        migrate()
        self.admin = create_user("admin", "admin-test-password", "admin")
        self.manager = create_user("manager", "manager-password", "manager", self.admin)
        self.other = create_user("other", "other-user-password", "manager", self.admin)
        self.client = self.client_for("manager", "manager-password")
        self.admin_client = self.client_for("admin", "admin-test-password")
        with connection(write=True) as db:
            self.order_id = create_order(db, example_response(), self.manager)

    def client_for(self, username, password):
        client = TestClient(app)
        self.addCleanup(client.close)
        response = client.post("/api/auth/login", json={"username": username, "password": password})
        self.assertEqual(response.status_code, 200, response.text)
        client.headers["X-CSRF-Token"] = response.json()["csrf_token"]
        return client

    def detail(self, client=None):
        return (client or self.client).get(f"/api/orders/{self.order_id}").json()

    def save(self, detail, client=None):
        return (client or self.client).put(f"/api/orders/{self.order_id}", json={"revision": detail["revision"], "lines": detail["decisions"]})

    def test_migrations_repeat_and_wal_enabled(self):
        migrate()
        migrate()
        with connection() as db:
            self.assertEqual(db.execute("PRAGMA user_version").fetchone()[0], 1)
            self.assertEqual(db.execute("PRAGMA journal_mode").fetchone()[0], "wal")
            self.assertEqual(db.execute("SELECT count(*) FROM orders").fetchone()[0], 1)

    def test_auth_protects_data_and_legacy_routes(self):
        with TestClient(app) as anonymous:
            for path in ("/api/meta", "/api/orders", "/api/jobs", "/api/users", "/api/auth/session"):
                self.assertEqual(anonymous.get(path).status_code, 401)
            self.assertEqual(anonymous.post("/api/recommend", json={}).status_code, 401)
            self.assertEqual(anonymous.get("/api/health").status_code, 200)

    def test_cookie_session_hash_csrf_logout_and_expiry(self):
        response = self.client.get("/api/auth/session")
        self.assertEqual(response.json()["user"]["role"], "manager")
        token = self.client.cookies.get(COOKIE_NAME)
        with connection() as db:
            self.assertIsNone(db.execute("SELECT 1 FROM sessions WHERE token_hash=?", (token,)).fetchone())
            password = db.execute("SELECT password_hash FROM users WHERE id=?", (self.manager["id"],)).fetchone()[0]
            self.assertTrue(password.startswith("$argon2id$"))
        headers = dict(self.client.headers)
        self.client.headers.pop("X-CSRF-Token")
        self.assertEqual(self.client.post("/api/auth/logout").status_code, 403)
        self.client.headers.update(headers)
        self.assertEqual(self.client.post("/api/auth/logout").status_code, 204)
        self.assertEqual(self.client.get("/api/auth/session").status_code, 401)
        self.client = self.client_for("manager", "manager-password")
        with connection(write=True) as db:
            db.execute("UPDATE sessions SET expires_at=0 WHERE user_id=?", (self.manager["id"],))
        self.assertEqual(self.client.get("/api/orders").status_code, 401)

    def test_request_logs_use_route_template_and_no_payload(self):
        with self.assertLogs("umytpa.requests", level="INFO") as captured:
            result = self.client.get("/api/orders/private-id?password=secret", headers={"X-Request-ID": "safe-request-id"})
        self.assertEqual(result.headers["X-Request-ID"], "safe-request-id")
        self.assertIn("route=/api/orders/{order_id}", " ".join(captured.output))
        self.assertNotIn("private-id", " ".join(captured.output))
        self.assertNotIn("secret", " ".join(captured.output))

    def test_production_cookie_secure_default_and_origin_rejected(self):
        with patch.dict(os.environ, {"APP_ENV": "production"}):
            with TestClient(app, base_url="https://testserver") as client:
                result = client.post("/api/auth/login", json={"username": "manager", "password": "manager-password"})
                cookie = result.headers["set-cookie"].lower()
                self.assertIn("secure", cookie)
                self.assertIn("httponly", cookie)
                self.assertIn("samesite=strict", cookie)
        attack = self.client.post("/api/auth/logout", headers={"Origin": "https://untrusted.example"})
        self.assertEqual(attack.status_code, 403)
        self.assertEqual(self.client.get("/api/auth/session").status_code, 200)

    def test_rate_limit_persists_across_clients(self):
        with TestClient(app) as anonymous:
            for _ in range(5):
                result = anonymous.post("/api/auth/login", json={"username": "missing", "password": "wrong"})
                self.assertEqual(result.status_code, 401)
        with TestClient(app) as another:
            result = another.post("/api/auth/login", json={"username": "missing", "password": "wrong"})
            self.assertEqual(result.status_code, 429)
            self.assertEqual(result.headers["Retry-After"], "900")

    def test_manager_cannot_access_another_order_or_admin(self):
        other = self.client_for("other", "other-user-password")
        self.assertEqual(other.get(f"/api/orders/{self.order_id}").status_code, 403)
        self.assertEqual(other.get(f"/api/orders/{self.order_id}/events").status_code, 403)
        self.assertEqual(other.get("/api/orders").json()["total"], 0)
        self.assertEqual(self.admin_client.get("/api/orders").json()["total"], 1)
        self.assertEqual(self.client.get("/api/users").status_code, 403)

    def test_user_deactivation_reset_revokes_sessions_last_admin_protected(self):
        result = self.admin_client.patch(f"/api/users/{self.manager['id']}", json={"active": False})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(self.client.get("/api/auth/session").status_code, 401)
        result = self.admin_client.patch(f"/api/users/{self.admin['id']}", json={"role": "manager"})
        self.assertEqual(result.status_code, 409)
        result = self.admin_client.patch(f"/api/users/{self.other['id']}", json={"password": "new-strong-password"})
        self.assertEqual(result.status_code, 200)
        with connection() as db:
            event = db.execute("SELECT changes FROM events WHERE action='user_updated' ORDER BY id DESC").fetchone()[0]
            self.assertNotIn("new-strong-password", event)
            self.assertIn("password_reset", event)

    def test_quantity_change_revokes_approval_and_conflict_is_atomic(self):
        original = self.detail()
        original["decisions"][0]["approved"] = True
        approved = self.save(original).json()
        self.assertTrue(approved["decisions"][0]["approved"])
        original["decisions"][0]["quantity"] = 25
        self.assertEqual(self.save(original).status_code, 409)
        approved["decisions"][0]["quantity"] = 25
        changed = self.save(approved).json()
        self.assertEqual(changed["decisions"][0]["quantity"], 25)
        self.assertFalse(changed["decisions"][0]["approved"])
        changed["decisions"][0]["approved"] = True
        saved = self.save(changed).json()
        self.assertTrue(saved["decisions"][0]["approved"])
        events = self.client.get(f"/api/orders/{self.order_id}/events").json()["items"]
        self.assertEqual(events[0]["action"], "decisions_updated")
        self.assertEqual(events[0]["actor_name"], "manager")
        self.assertEqual(self.detail()["decisions"], saved["decisions"])

    def test_exact_line_set_and_invalid_quantity_rollback(self):
        for mode in ("duplicate", "missing", "unknown", "minimum", "multiple", "zero_approval", "bool"):
            data = self.detail()
            if mode == "duplicate":
                data["decisions"][1]["line_id"] = data["decisions"][0]["line_id"]
            elif mode == "missing":
                data["decisions"].pop()
            elif mode == "unknown":
                data["decisions"][0]["line_id"] = "unknown"
            elif mode == "minimum":
                data["decisions"][1]["quantity"] = 5
            elif mode == "multiple":
                data["decisions"][1]["quantity"] = 11
            elif mode == "bool":
                data["decisions"][0]["quantity"] = True
            else:
                data["decisions"][0].update(quantity=0, approved=True)
            with self.subTest(mode=mode):
                self.assertEqual(self.save(data).status_code, 422)
                self.assertEqual(self.detail()["revision"], 1)

    def test_archive_restore_and_export_no_recalculation(self):
        data = self.detail()
        data["decisions"][0]["approved"] = True
        data = self.save(data).json()
        with patch("app.jobs.get_data_source", side_effect=AssertionError("no reload")):
            result = self.client.post(f"/api/orders/{self.order_id}/export", json={"revision": data["revision"], "approved_only": True})
        self.assertEqual(result.status_code, 200)
        book = load_workbook(io.BytesIO(result.content), data_only=True)
        self.addCleanup(book.close)
        self.assertEqual(book["Заказ поставщикам"].max_row, 2)
        archived = self.client.patch(f"/api/orders/{self.order_id}", json={"revision": data["revision"], "archived": True, "title": "Сентябрь"}).json()
        self.assertEqual(self.client.get("/api/orders").json()["total"], 0)
        self.assertEqual(self.client.get("/api/orders?archived=true").json()["total"], 1)
        self.assertEqual(self.save(archived).status_code, 409)
        restored = self.client.patch(f"/api/orders/{self.order_id}", json={"revision": archived["revision"], "archived": False})
        self.assertEqual(restored.status_code, 200)
        self.assertEqual(restored.json()["title"], "Сентябрь")

    def test_empty_calculation_can_be_saved_but_cannot_export(self):
        empty = example_response().model_copy(update={"groups": [], "sku_count": 0})
        with connection(write=True) as db:
            order_id = create_order(db, empty, self.manager)
        result = self.client.put(f"/api/orders/{order_id}", json={"revision": 1, "lines": []})
        self.assertEqual(result.status_code, 200)
        result = self.client.post(f"/api/orders/{order_id}/export", json={"revision": 1, "approved_only": False})
        self.assertEqual(result.status_code, 422)

    def test_metadata_sanitizes_unix_and_windows_paths(self):
        data = public_metadata({"files": [{"file": "C:\\secret\\IEK\\data.xlsx", "sheet": "Лист1", "path": "/secret"}, {"file": "/srv/secret/data2.xlsx"}],
                                "rows": {"catalog": 4}, "OPENAI_API_KEY": "secret", "root": "/secret"})
        self.assertEqual(data["files"][0]["file"], "data.xlsx")
        self.assertEqual(data["files"][1]["file"], "data2.xlsx")
        self.assertNotIn("secret", json.dumps(data))

    def test_queue_requires_worker_idempotency_and_access(self):
        headers = {"Prefer": "respond-async", "Idempotency-Key": "first-key"}
        self.assertEqual(self.client.post("/api/recommend", json={}, headers=headers).status_code, 503)
        self.assertTrue(register("worker-one"))
        first = self.client.post("/api/recommend", json={}, headers=headers)
        self.assertEqual(first.status_code, 202, first.text)
        repeat = self.client.post("/api/recommend", json={}, headers=headers)
        self.assertEqual(first.json(), repeat.json())
        different = self.client.post("/api/recommend", json={"review_period_days": 20}, headers=headers)
        self.assertEqual(different.status_code, 409)
        other = self.client_for("other", "other-user-password")
        job_id = first.json()["job_id"]
        self.assertEqual(other.get(f"/api/jobs/{job_id}").status_code, 403)
        cancelled = self.client.post(f"/api/jobs/{job_id}/cancel")
        self.assertEqual(cancelled.json()["status"], "cancelled")
        self.assertEqual(self.client.get("/api/jobs").json()["items"][0]["status"], "cancelled")

    def test_job_completes_once_and_cancelled_job_cannot_publish(self):
        register("worker-one")
        command = jobs.enqueue(__import__("app.schemas", fromlist=["RecommendRequest"]).RecommendRequest(), self.manager, "job-key")
        claimed = jobs.claim("worker-one")
        self.assertEqual(claimed, command["job_id"])
        with patch("app.jobs.get_data_source") as source, patch("app.jobs.generate_recommendations", return_value=example_response()):
            source.return_value.load.return_value = object()
            jobs.execute_job(claimed, "worker-one")
            jobs.execute_job(claimed, "worker-one")
        completed = jobs.get_job(claimed, self.manager)
        self.assertEqual(completed["status"], "completed")
        self.assertIsNotNone(completed["order_id"])
        next_job = jobs.enqueue(__import__("app.schemas", fromlist=["RecommendRequest"]).RecommendRequest(), self.manager, "next-key")
        jobs.claim("worker-one")
        def cancel_during_calculation(*_, **__):
            jobs.cancel(next_job["job_id"], self.manager)
            return example_response()
        with patch("app.jobs.get_data_source"), patch("app.jobs.generate_recommendations", side_effect=cancel_during_calculation):
            jobs.execute_job(next_job["job_id"], "worker-one")
        cancelled = jobs.get_job(next_job["job_id"], self.manager)
        self.assertEqual(cancelled["status"], "cancelled")
        self.assertIsNone(cancelled["order_id"])
        self.assertEqual(self.client.get("/api/orders").json()["total"], 2)

    def test_worker_recovery_fails_running_keeps_queue(self):
        self.assertTrue(register("worker-one"))
        from app.schemas import RecommendRequest
        first = jobs.enqueue(RecommendRequest(), self.manager, "first")
        second = jobs.enqueue(RecommendRequest(), self.other, "second")
        jobs.claim("worker-one")
        self.assertFalse(register("worker-two"))
        with connection(write=True) as db:
            db.execute("UPDATE worker_state SET heartbeat=0")
        self.assertTrue(register("worker-two"))
        self.assertFalse(heartbeat("worker-one"))
        self.assertIsNone(jobs.claim("worker-one"))
        self.assertEqual(jobs.get_job(first["job_id"], self.manager)["status"], "failed")
        self.assertEqual(jobs.get_job(second["job_id"], self.other)["status"], "queued")

    def test_failed_job_uses_safe_message(self):
        register("worker-one")
        from app.schemas import RecommendRequest
        job = jobs.enqueue(RecommendRequest(), self.manager, "failure")
        jobs.claim("worker-one")
        with patch("app.jobs.get_data_source", side_effect=RuntimeError("/private/secret credentials")):
            jobs.execute_job(job["job_id"], "worker-one")
        result = jobs.get_job(job["job_id"], self.manager)
        self.assertEqual(result["status"], "failed")
        self.assertNotIn("secret", result["error"])

    def test_llm_budget_fallback_never_calls_after_deadline(self):
        line = example_response().groups[0].lines[0]
        with patch("app.core.explain.get_settings") as settings, patch("openai.OpenAI") as client:
            settings.return_value.llm_enabled = True
            with explanation_budget(-1):
                text = build_explanation(line.name, line.rationale, 20, "high", True)
            self.assertIn("20", text)
            client.assert_not_called()


class MigrationTests(unittest.TestCase):
    def test_legacy_snapshots_import_once_archived_without_approvals(self):
        with tempfile.TemporaryDirectory() as folder, patch.dict(os.environ, {"APP_DB_PATH": folder + "/db.sqlite3", "ORDER_DB_PATH": folder + "/db.sqlite3"}):
            snapshot = save_snapshot(example_response())
            migrate()
            admin = create_user("first-admin", "long-test-password", "admin")
            create_user("second-admin", "long-test-password", "admin")
            with connection() as db:
                rows = db.execute("SELECT * FROM orders").fetchall()
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["owner_id"], admin["id"])
                self.assertEqual(rows[0]["archived"], 1)
                self.assertEqual(rows[0]["calculation_id"], snapshot.calculation_id)
                self.assertEqual(db.execute("SELECT sum(approved) FROM decisions").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
