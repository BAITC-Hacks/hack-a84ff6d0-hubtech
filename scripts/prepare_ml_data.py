#!/usr/bin/env python3
"""Подготовка шести IEK-выгрузок и признаков месячного прогноза.

python scripts/prepare_ml_data.py --as-of 2026-09-23
Зависимости: pandas, openpyxl. Источники не изменяются; модель не обучается.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from clean_monthly_data import ROOT, SALES_FILE, STOCK_FILE, clean_text, parse_month, parse_quantity

FILES = {
    "sales": SALES_FILE, "stock": STOCK_FILE,
    "transactions": "Динамика продаж_2025-2026.xlsx",
    "moq": "MOQ  ИЭК.xlsx", "transit": "Путь ИЭК 22.09.2026.xlsx",
    "seasonality": "Сезонность ИЭК.xlsx",
}
LAGS = (1, 2, 3, 6, 12)
FEATURES = [
    "sku", "unit", "month_number", "days_in_month",
    *[f"lag_{lag}" for lag in LAGS], "mean_3", "mean_6", "std_6",
    "observations_3", "observations_6", "missing_share_6", "zero_share_6",
    "history_observed_months", "history_calendar_months", "opening_stock_lag_1",
]
ISSUE_COLUMNS = ["source", "row", "column", "sku", "issue", "value"]


@contextmanager
def table(path, anchor):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        rows = enumerate(sheet.iter_rows(values_only=True), 1)
        for row_number, row in rows:
            headers = [clean_text(v).casefold() for v in row]
            if anchor.casefold() in headers:
                yield headers, rows, row
                return
            if row_number >= 20:
                break
        raise ValueError(f"{path.name}: не найдена шапка с полем {anchor}")
    finally:
        workbook.close()


def value(row, index):
    return row[index] if index < len(row) else None


def issue(issues, source, row, column, sku, kind, raw):
    issues.append(dict(source=source, row=row, column=column, sku=sku, issue=kind,
                       value="" if raw is None else str(raw)))


def quantity(raw, issues, source, row, column, sku):
    try:
        parsed = parse_quantity(raw)
    except ValueError:
        issue(issues, source, row, column, sku, "invalid_quantity", raw)
        return np.nan, "invalid"
    if parsed is None:
        return np.nan, "blank"
    if parsed < 0:
        issue(issues, source, row, column, sku, "negative_quantity_preserved", raw)
        return parsed, "negative"
    return parsed, "observed"


def code(raw):
    if raw is None or clean_text(raw).casefold() in {"", "итого", "всего"}:
        return ""
    if not isinstance(raw, str):
        raise ValueError(f"Код {raw!r} хранится числом: проверьте ведущие нули в источнике")
    return clean_text(raw)


def catalog_add(catalog, sku, name="", unit="", supplier_sku=""):
    record = catalog.setdefault(sku, dict(sku=sku, product_name="", unit="", supplier_sku=""))
    for key, raw in [("product_name", name), ("unit", unit), ("supplier_sku", supplier_sku)]:
        text = clean_text(raw)
        if key in {"unit", "supplier_sku"} and text and record[key] and record[key] != text:
            label = "единицы" if key == "unit" else "артикулы поставщика"
            raise ValueError(f"{sku}: несовместимые {label} {record[key]!r} и {text!r}")
        if text and not record[key]:
            record[key] = text


def read_matrix(path, kind, catalog, issues):
    records, seen = [], {}
    with table(path, "Номенклатура.Код") as (headers, rows, raw_headers):
        ci, ni = headers.index("номенклатура.код"), headers.index("номенклатура")
        ui = headers.index("ед.") if "ед." in headers else None
        months = {i: pd.Timestamp(parse_month(h)) for i, h in enumerate(raw_headers) if parse_month(h)}
        if not months or len(set(months.values())) != len(months):
            raise ValueError(f"{path.name}: месяцы отсутствуют или повторяются")
        for row_number, row in rows:
            if clean_text(value(row, ni)).casefold() in {"итого", "всего"}:
                continue
            sku = code(value(row, ci))
            if not sku:
                if clean_text(value(row, ni)):
                    issue(issues, kind, row_number, "Номенклатура.Код", "", "missing_code", value(row, ni))
                continue
            name, unit = clean_text(value(row, ni)), clean_text(value(row, ui)) if ui is not None else ""
            signature = (name, unit, tuple(value(row, i) for i in months))
            if sku in seen:
                if signature != seen[sku]:
                    raise ValueError(f"{path.name}: противоречивые строки для {sku}")
                issue(issues, kind, row_number, "Номенклатура.Код", sku, "duplicate_removed", "")
                continue
            seen[sku] = signature
            catalog_add(catalog, sku, name, unit)
            for col, month in months.items():
                qty, status = quantity(value(row, col), issues, kind, row_number, headers[col], sku)
                records.append(dict(sku=sku, month=month, **{
                    f"{kind}_qty_raw": qty, f"{kind}_status": status,
                }))
    if not records:
        raise ValueError(f"{path.name}: нет данных")
    return pd.DataFrame(records)


def read_transactions(path, catalog, issues):
    records = []
    names = ["дата", "номер", "документ", "код", "номенклатура", "ед.", "склад", "количество"]
    with table(path, "Дата") as (headers, rows, _):
        columns = [headers.index(name) for name in names]
        for row_number, row in rows:
            raw_date, number, doc, raw_code, name, unit, warehouse, raw_qty = [value(row, i) for i in columns]
            sku = code(raw_code)
            if not sku:
                continue
            try:
                date = pd.Timestamp(datetime.strptime(raw_date, "%d.%m.%Y %H:%M:%S")) if isinstance(raw_date, str) else pd.Timestamp(raw_date)
            except (ValueError, TypeError):
                date = pd.NaT
            if pd.isna(date):
                issue(issues, "transactions", row_number, "Дата", sku, "invalid_date", raw_date)
            qty, status = quantity(raw_qty, issues, "transactions", row_number, "Количество", sku)
            doc = clean_text(doc)
            doc_type = next((kind for kind in ("Расходная накладная", "Приходная накладная", "Заказ покупателя")
                             if doc.startswith(kind)), "unknown")
            catalog_add(catalog, sku, name, unit)
            records.append(dict(sku=sku, date=date, document_number=clean_text(number), document_type=doc_type,
                                document=doc, warehouse=clean_text(warehouse), unit=clean_text(unit),
                                qty_raw=qty, quantity_status=status, source_row=row_number,
                                is_positive_outgoing=doc_type == "Расходная накладная" and qty > 0))
    frame = pd.DataFrame(records)
    if frame.empty:
        raise ValueError("Нет транзакций")
    return frame  # Одинаковые строки могут быть позициями документа: не удаляем автоматически.


def read_moq(path, catalog, issues):
    records = {}
    with table(path, "Код 1с") as (headers, rows, _):
        ci, ai, ni, qi = [headers.index(h) for h in ("код 1с", "артикул поставщика", "наименование", "мин. разр. к отгр.")]
        for rn, row in rows:
            sku = code(value(row, ci))
            if not sku:
                continue
            article = clean_text(value(row, ai))
            catalog_add(catalog, sku, value(row, ni), supplier_sku=article)
            qty, status = quantity(value(row, qi), issues, "moq", rn, headers[qi], sku)
            minimum = qty if np.isfinite(qty) and qty > 0 else np.nan
            if status == "observed" and not qty > 0:
                issue(issues, "moq", rn, headers[qi], sku, "invalid_minimum", qty)
            current = dict(sku=sku, supplier_sku=article, min_order_qty=minimum,
                           moq_known=bool(np.isfinite(minimum)))
            if sku in records:
                old = records[sku]
                same_qty = (pd.isna(old["min_order_qty"]) and pd.isna(minimum)) or old["min_order_qty"] == minimum
                if not same_qty or article != old["supplier_sku"]:
                    raise ValueError(f"Конфликт MOQ для {sku}; автоматический выбор запрещён")
                issue(issues, "moq", rn, "Код 1с", sku, "duplicate_constraint_removed", article)
            else:
                records[sku] = current
    return pd.DataFrame(records.values(), columns=["sku", "supplier_sku", "min_order_qty", "moq_known"])


def read_transit(path, catalog, issues):
    match = re.search(r"(\d{2}\.\d{2}\.\d{4})", path.stem)
    if not match:
        raise ValueError("Дата снимка пути должна быть в имени файла")
    snapshot_date = pd.Timestamp(datetime.strptime(match[1], "%d.%m.%Y"))
    records, seen = [], set()
    with table(path, "Код 1с") as (headers, rows, _):
        ci, ai, ni = [headers.index(h) for h in ("код 1с", "артикул иэк", "наименование")]
        arrivals = {}
        for i, header in enumerate(headers):
            eta = re.search(r"поступление\s+до\s+(\d{2}\.\d{2}\.\d{4})", header)
            if eta:
                arrivals[i] = pd.Timestamp(datetime.strptime(eta[1], "%d.%m.%Y"))
        if not arrivals:
            raise ValueError("В пути отсутствуют даты поступления")
        for rn, row in rows:
            raw_code = value(row, ci)
            # Только две подтверждённые формы заметок: коды могут быть числом
            # или текстом. Реальные товары с кодом 0/1 не исключаются.
            article = clean_text(value(row, ai))
            note_shape = (not clean_text(value(row, ni))
                          and ((clean_text(raw_code) == "0" and article.casefold().startswith("расширение "))
                               or (clean_text(raw_code) == "1" and article == "1")))
            try:
                no_receipts = all(parse_quantity(value(row, col)) in (None, 0) for col in arrivals)
            except ValueError:
                no_receipts = False
            if note_shape and no_receipts:
                issue(issues, "transit", rn, "Код 1с", "", "non_product_note_row_excluded", row)
                continue
            sku = code(raw_code)
            if not sku:
                continue
            catalog_add(catalog, sku, value(row, ni), supplier_sku=value(row, ai))
            for col, eta in arrivals.items():
                raw = value(row, col)
                if raw is None or clean_text(raw) == "":
                    continue
                key = (sku, col)
                if key in seen:
                    raise ValueError(f"Путь: повтор {sku} в поставке {headers[col]}; нужна ручная сверка")
                seen.add(key)
                qty, status = quantity(raw, issues, "transit", rn, headers[col], sku)
                records.append(dict(sku=sku, shipment=headers[col], eta=eta, qty_raw=qty,
                                    quantity_status=status, source_as_of=snapshot_date,
                                    unit_conversion_confirmed=False))
    return pd.DataFrame(records, columns=["sku", "shipment", "eta", "qty_raw", "quantity_status", "source_as_of", "unit_conversion_confirmed"])


def read_seasonality(path, issues):
    records = []
    month_names = {name: i for i, name in enumerate("янв фев мар апр май июн июл авг сен окт ноя дек".split(), 1)}
    with table(path, "СЕЗОННОСТЬ") as (headers, rows, _):
        ci, mi = headers.index("сезонность"), headers.index("месяц")
        for rn, row in rows:
            month = clean_text(value(row, mi)).casefold().rstrip(".")
            if month not in month_names:
                if records:
                    break
                continue
            factor, _ = quantity(value(row, ci), issues, "seasonality", rn, "СЕЗОННОСТЬ", "")
            records.append(dict(month_number=month_names[month], provided_factor=factor,
                                usable_for_historical_training=False))
    if len(records) != 12:
        raise ValueError("В сезонности ожидается 12 коэффициентов")
    return pd.DataFrame(records)


def build_features(sales, stock, products, cutoff, blank_sales="missing", negative_sales="missing", min_history=3):
    """Все признаки строки месяца T зависят только от месяцев < T."""
    prediction_month = pd.Timestamp(cutoff).to_period("M").to_timestamp()
    start = min(sales.month.min(), stock.month.min())
    if start >= prediction_month:
        raise ValueError("Недостаточно завершённых месяцев до даты расчёта")
    index = pd.MultiIndex.from_product([products.sku, pd.date_range(start, prediction_month, freq="MS")], names=["sku", "month"])
    panel = index.to_frame(index=False).merge(sales, how="left", on=["sku", "month"], validate="one_to_one")
    panel = panel.merge(stock, how="left", on=["sku", "month"], validate="one_to_one")
    panel = panel.merge(products[["sku", "unit"]], how="left", on="sku", validate="many_to_one")
    panel = panel.sort_values(["sku", "month"]).reset_index(drop=True)
    panel["is_complete_month"] = panel.month < prediction_month
    panel["sales_status"] = panel.sales_status.fillna("not_in_source")
    panel["stock_status"] = panel.stock_status.fillna("not_in_source")
    qty = panel.sales_qty_raw.copy()
    qty.loc[panel.sales_status.isin(["invalid", "not_in_source"])] = np.nan
    if blank_sales == "zero":
        qty.loc[panel.sales_status.eq("blank")] = 0.0
    qty.loc[qty < 0] = 0.0 if negative_sales == "zero" else np.nan
    qty.loc[~panel.is_complete_month] = np.nan
    panel["target_qty"] = qty
    # Начальный остаток предыдущего месяца можно знать в момент прогноза T.
    stock_qty = panel.stock_qty_raw.where(panel.stock_qty_raw.ge(0) & panel.is_complete_month)
    panel["stock_qty_model"] = stock_qty
    grouped = panel.groupby("sku", sort=False)
    for lag in LAGS:
        panel[f"lag_{lag}"] = grouped.target_qty.shift(lag)
    panel["opening_stock_lag_1"] = grouped.stock_qty_model.shift(1)
    panel["history_calendar_months"] = grouped.cumcount()
    panel["history_observed_months"] = grouped.target_qty.transform(lambda s: s.notna().cumsum() - s.notna().astype(int))
    lagged = panel.groupby("sku", sort=False).lag_1
    for window in (3, 6):
        panel[f"mean_{window}"] = lagged.transform(lambda s: s.rolling(window, min_periods=1).mean())
        panel[f"observations_{window}"] = lagged.transform(lambda s: s.rolling(window, min_periods=1).count())
    panel["std_6"] = lagged.transform(lambda s: s.rolling(6, min_periods=2).std())
    panel["missing_share_6"] = 1 - panel.observations_6 / panel.history_calendar_months.clip(upper=6).replace(0, np.nan)
    zeros = lagged.transform(lambda s: s.eq(0).astype(int).rolling(6, min_periods=1).sum())
    panel["zero_share_6"] = zeros / panel.observations_6.replace(0, np.nan)
    panel["month_number"] = panel.month.dt.month
    panel["days_in_month"] = panel.month.dt.days_in_month
    panel["history_sufficient"] = panel.history_observed_months >= min_history
    train = panel[panel.is_complete_month & panel.target_qty.notna() & panel.history_sufficient].copy()
    # Общие временные границы для всех SKU: 6 месяцев validation и 2 test.
    test_start = prediction_month - pd.DateOffset(months=2)
    validation_start = test_start - pd.DateOffset(months=6)
    train["split"] = np.select([train.month < validation_start, train.month < test_start], ["train", "validation"], default="test")
    train = train.rename(columns={"month": "target_month"})[["target_month", *FEATURES, "target_qty", "split"]]
    predict = panel[panel.month.eq(prediction_month)].rename(columns={"month": "target_month"})
    predict = predict[["target_month", *FEATURES, "history_sufficient"]]
    return panel, train, predict


def reconcile(sales, transactions, cutoff):
    complete = pd.Timestamp(cutoff).to_period("M").to_timestamp()
    selected = transactions[transactions.document_type.eq("Расходная накладная") & transactions.date.lt(complete)].copy()
    selected["month"] = selected.date.dt.to_period("M").dt.to_timestamp()
    selected["positive_qty"] = selected.qty_raw.clip(lower=0)
    groups = selected.groupby(["sku", "month"])
    totals = groups[["qty_raw", "positive_qty"]].sum(min_count=1)
    totals["transaction_unknown_qty_rows"] = groups.size() - groups.qty_raw.count()
    totals = totals.reset_index().rename(columns={"qty_raw": "transactions_signed_qty", "positive_qty": "transactions_positive_qty"})
    result = sales[sales.month.lt(complete)].merge(totals, how="outer", on=["sku", "month"], validate="one_to_one")
    valid = (result.sales_qty_raw.notna() & result.transactions_signed_qty.notna()
             & result.transaction_unknown_qty_rows.eq(0))
    result["comparable"] = valid
    result["matches_signed"] = valid & np.isclose(result.sales_qty_raw, result.transactions_signed_qty, rtol=1e-8, atol=1e-6)
    result["matches_positive"] = valid & np.isclose(result.sales_qty_raw, result.transactions_positive_qty, rtol=1e-8, atol=1e-6)
    result["difference_signed"] = result.transactions_signed_qty - result.sales_qty_raw
    return result


def prepare(input_dir, output_dir, as_of, blank_sales="missing", negative_sales="missing", min_history=3):
    if blank_sales not in {"missing", "zero"} or negative_sales not in {"missing", "zero"} or min_history < 1:
        raise ValueError("Некорректные правила очистки или min_history")
    as_of = pd.Timestamp(as_of).normalize()
    if pd.isna(as_of):
        raise ValueError("Некорректная дата расчёта")
    paths = {kind: input_dir / name for kind, name in FILES.items()}
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(path)
    catalog, issues = {}, []
    stock = read_matrix(paths["stock"], "stock", catalog, issues)
    sales = read_matrix(paths["sales"], "sales", catalog, issues)
    transactions = read_transactions(paths["transactions"], catalog, issues)
    moq = read_moq(paths["moq"], catalog, issues)
    transit = read_transit(paths["transit"], catalog, issues)
    seasonality = read_seasonality(paths["seasonality"], issues)
    products = pd.DataFrame(catalog.values()).sort_values("sku").reset_index(drop=True)
    latest_transaction = transactions.date.max()
    if pd.isna(latest_transaction):
        raise ValueError("Нет корректных дат транзакций для определения границы данных")
    # Смена --as-of не должна превращать неполный месяц старого файла в полный.
    cutoff = min(as_of, latest_transaction.normalize() + pd.Timedelta(days=1))
    panel, train, predict = build_features(sales, stock, products, cutoff, blank_sales, negative_sales, min_history)
    comparison = reconcile(sales, transactions, cutoff)
    warnings = [
        "Пропуски и возвраты требуют подтверждения партнёром. Исходные количества сохранены.",
        "MOQ и путь — отдельные входы расчёта закупок; их нет среди ML-признаков.",
        "Готовая сезонность содержит данные 2026 и прогнозные компоненты; в признаки она не включена.",
        "Дни отсутствия товара, ID клиента и актуальный остаток IEK отсутствуют; они не выдумываются.",
        "Месячные отчёты не содержат склада. Соответствие области отчётов складу транзакций не подтверждено.",
        "Прогнозный месяц следует за завершённой историей. При старом снимке он может быть раньше --as-of.",
        "Пустые продажи не доказывают нулевой спрос; --blank-sales zero — явное допущение, включая месяцы до первых продаж.",
        "Строки raw-таблиц могут быть позже исторического --as-of: используйте только явный список признаков.",
        "Единицы пути/закупки могут отличаться от складских (например, бухта/метр); пересчёт не подтверждён.",
        "Большие продажи не удалены: нет подтверждённой разметки опта, источники объёмов расходятся.",
        "Набор охватывает только IEK. Systeme Electric не входит в этот конвейер.",
        "Цели с неизвестными продажами исключены: метрики по ним не доказывают качество прогноза полного спроса.",
    ]
    report = {
        "as_of": as_of.date().isoformat(), "data_cutoff_exclusive": cutoff.date().isoformat(),
        "prediction_month": predict.target_month.iloc[0].date().isoformat(),
        "policies": {"blank_sales": blank_sales, "negative_sales": negative_sales, "min_history": min_history},
        "feature_columns": FEATURES, "categorical_features": ["sku", "unit", "month_number"],
        "target_column": "target_qty", "time_column": "target_month", "group_column": "sku",
        "scope": {"supplier": "IEK", "monthly_warehouse_binding_confirmed": False,
                  "transaction_warehouses": sorted(transactions.warehouse.dropna().unique().tolist()),
                  "runtime_csv_source_compatible": False},
        "target_profile": {"training_targets": len(train),
                           "positive_targets": int(train.target_qty.gt(0).sum()),
                           "zero_targets": int(train.target_qty.eq(0).sum()),
                           "observed_target_selection": True,
                           "total_demand_validated": False},
        "counts": {"products": len(products), "monthly_rows": len(panel), "training_rows": len(train),
                   "prediction_rows": len(predict), "prediction_with_sufficient_history": int(predict.history_sufficient.sum()),
                   "transactions": len(transactions), "moq_rows": len(moq), "transit_receipts": len(transit),
                   "comparable_sku_months": int(comparison.comparable.sum()),
                   "mismatched_signed_sku_months": int((comparison.comparable & ~comparison.matches_signed).sum())},
        "splits": {name: {"rows": len(group), "start": group.target_month.min().date().isoformat(),
                          "end": group.target_month.max().date().isoformat()} for name, group in train.groupby("split")},
        "quantity_statuses": {"sales": sales.sales_status.value_counts().to_dict(), "stock": stock.stock_status.value_counts().to_dict(),
                              "transactions": transactions.quantity_status.value_counts().to_dict()},
        "issues": dict(Counter(item["issue"] for item in issues)),
        "unknown_units": products.loc[products.unit.eq(""), "sku"].tolist(),
        "warnings": warnings,
        "sources": {kind: {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for kind, path in paths.items()},
    }
    if train.empty:
        raise ValueError("Нет пригодных строк обучения при заданных правилах и min_history")
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = dict(products=products, monthly_panel=panel, transactions=transactions, moq=moq,
                   in_transit=transit, provided_seasonality=seasonality, reconciliation=comparison,
                   training_data=train, prediction_features=predict,
                   quality_issues=pd.DataFrame(issues, columns=ISSUE_COLUMNS))
    for name, frame in outputs.items():
        frame.to_csv(output_dir / f"{name}.csv", index=False, encoding="utf-8-sig", na_rep="")
    (output_dir / "preparation_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=ROOT / "IEK")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data" / "ml" / "iek")
    parser.add_argument("--as-of", default="2026-09-23", help="Дата расчёта; по умолчанию дата текущих выгрузок")
    parser.add_argument("--blank-sales", choices=["missing", "zero"], default="missing")
    parser.add_argument("--negative-sales", choices=["missing", "zero"], default="missing")
    parser.add_argument("--min-history", type=int, default=3, help="Минимум известных месяцев в прошлом")
    args = parser.parse_args()
    try:
        report = prepare(args.input_dir, args.output_dir, args.as_of, args.blank_sales, args.negative_sales, args.min_history)
    except (ValueError, KeyError, OSError) as error:
        parser.exit(1, f"Ошибка подготовки: {error}\n")
    print(json.dumps({"counts": report["counts"], "splits": report["splits"]}, ensure_ascii=False, indent=2))
    print(f"Готово: {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
