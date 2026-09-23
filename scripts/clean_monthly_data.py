#!/usr/bin/env python3
"""Очистка месячных продаж и начальных остатков IEK без изменения источников.

Установка: pip install pandas openpyxl
Запуск: python scripts/clean_monthly_data.py
"""

import argparse
import json
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
STOCK_FILE = "Ежемесячные остатки продукции за последние 2 года  ИЭК.xlsx"
SALES_FILE = "Ежемесячные продажи в количественном выражении за последние 2 года.xlsx"
MONTHS = {
    "янв": 1, "февр": 2, "март": 3, "апр": 4, "май": 5, "июнь": 6,
    "июль": 7, "авг": 8, "сент": 9, "окт": 10, "нояб": 11, "дек": 12,
}
ISSUE_COLUMNS = ["source", "row", "column", "product_code", "issue", "value"]


def clean_text(value):
    """Убираем лишние пробелы; коды остаются строками, включая 0 и _."""
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split())


def parse_month(value):
    if isinstance(value, (date, datetime)):
        return date(value.year, value.month, 1).isoformat()
    match = re.fullmatch(r"([а-я]+)\.?\s+(\d{4})", clean_text(value).lower())
    if match and match[1] in MONTHS:
        return date(int(match[2]), MONTHS[match[1]], 1).isoformat()
    return None


def parse_quantity(value):
    """Пусто != 0. Отрицательные количества сохраняем для проверки."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise ValueError("Логическое значение вместо количества")
    text = str(value).replace("\xa0", "").replace("\u202f", "").replace(" ", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:[eE][+-]?\d+)?", text):
        raise ValueError("Некорректное количество")
    number = float(text.replace(",", "."))
    if not math.isfinite(number):
        raise ValueError("Количество должно быть конечным числом")
    return number


def read_monthly(path, kind, sheet_name, issues):
    """Читаем двух-/трёхстрочную шапку и разворачиваем месяцы в строки."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        header_row, headers = None, None
        for header_number, cells in enumerate(sheet.iter_rows(max_row=20, values_only=True), 1):
            labels = [clean_text(value) for value in cells]
            if "Номенклатура.Код" in labels:
                headers = labels
                header_row = cells
                break
        if headers is None:
            raise ValueError(f"{path.name}: не найден заголовок Номенклатура.Код")

        code_col = headers.index("Номенклатура.Код")
        name_col = headers.index("Номенклатура")
        unit_col = headers.index("Ед.") if "Ед." in headers else None
        if kind == "stock" and unit_col is None:
            raise ValueError(f"{path.name}: не найден столбец Ед.")
        month_cols = {i: parse_month(value) for i, value in enumerate(header_row)
                      if parse_month(value)}
        if not month_cols or len(set(month_cols.values())) != len(month_cols):
            raise ValueError(f"{path.name}: месяцы отсутствуют или повторяются")

        stats = Counter()
        seen, products, records = {}, [], []

        def note(row, col, code, issue, value):
            issues.append(dict(source=kind, row=row, column=get_column_letter(col + 1),
                               product_code=code, issue=issue, value=str(value)))

        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_number + 1, values_only=True), header_number + 1
        ):
            name, code = clean_text(row[name_col]), clean_text(row[code_col])
            if name.casefold() in {"итого", "всего"}:
                stats["total_rows_removed"] += 1
                continue
            if not name and not code:
                # Подзаголовки «Количество», «нач. остаток» и пустые строки.
                values = {clean_text(v).casefold() for v in row if clean_text(v)}
                if values <= {"количество", "нач. остаток"}:
                    stats["service_rows_removed"] += 1
                    continue
            if not code:
                note(row_number, code_col, "", "missing_product_code_row_excluded", name)
                stats["rows_without_code"] += 1
                continue
            if not isinstance(row[code_col], str):
                raise ValueError(f"{path.name}, строка {row_number}: код должен быть текстом; "
                                 "проверьте ведущие нули в исходном файле")
            unit = clean_text(row[unit_col]) if unit_col is not None else ""
            signature = (name, unit, tuple(row[i] for i in month_cols))
            if code in seen:
                previous_row, previous_signature = seen[code]
                if signature != previous_signature:
                    raise ValueError(f"{path.name}: разные записи для кода {code}, строки "
                                     f"{previous_row} и {row_number}. Автосуммирование запрещено.")
                note(row_number, code_col, code, "exact_duplicate_removed", previous_row)
                stats["exact_duplicates_removed"] += 1
                continue
            seen[code] = (row_number, signature)
            products.append(dict(product_code=code, name=name, unit=unit))
            for col, month in month_cols.items():
                try:
                    quantity = parse_quantity(row[col])
                except ValueError:
                    note(row_number, col, code, "invalid_quantity_kept_as_missing", row[col])
                    stats["invalid_quantity_cells"] += 1
                    quantity = None
                else:
                    if quantity is None:
                        stats["blank_quantity_cells"] += 1
                    elif quantity < 0:
                        note(row_number, col, code, "negative_quantity_preserved", row[col])
                        stats["negative_quantity_cells"] += 1
                records.append(dict(product_code=code, month=month,
                                    **{f"{kind}_qty": quantity, f"{kind}_row_present": True}))
        if not products:
            raise ValueError(f"{path.name}: нет товарных строк")
        return (
            pd.DataFrame(records), pd.DataFrame(products).set_index("product_code"),
            dict(stats, file=str(path.resolve()), sheet=sheet.title, products=len(products),
                 months=sorted(month_cols.values()), output_rows=len(records)),
        )
    finally:
        workbook.close()


def clean_datasets(stock_path, sales_path, output_dir, sheet_name=None):
    issues = []
    stock, stock_products, stock_stats = read_monthly(stock_path, "stock", sheet_name, issues)
    sales, sales_products, sales_stats = read_monthly(sales_path, "sales", sheet_name, issues)
    products = stock_products.replace("", pd.NA).combine_first(sales_products.replace("", pd.NA))
    for code in stock_products.index.intersection(sales_products.index):
        stock_unit, sales_unit = stock_products.at[code, "unit"], sales_products.at[code, "unit"]
        if stock_unit and sales_unit and stock_unit != sales_unit:
            raise ValueError(f"{code}: единицы продаж и остатков не совпадают: {sales_unit} / {stock_unit}")
        left, right = stock_products.at[code, "name"], sales_products.at[code, "name"]
        if left and right and left != right:
            issues.append(dict(source="both", row="", column="Номенклатура", product_code=code,
                               issue="name_mismatch_stock_name_used", value=f"{left} | {right}"))

    result = sales.merge(stock, on=["product_code", "month"], how="outer", validate="one_to_one")
    result = result.merge(products.reset_index(), on="product_code", how="left", validate="many_to_one")
    for kind in ("sales", "stock"):
        result[f"{kind}_row_present"] = result[f"{kind}_row_present"].eq(True)
    result = result.rename(columns={"stock_qty": "opening_stock_qty", "name": "product_name"})
    result = result[["product_code", "product_name", "unit", "month", "sales_qty",
                     "opening_stock_qty", "sales_row_present", "stock_row_present"]]
    result = result.sort_values(["product_code", "month"]).reset_index(drop=True)
    report = {
        "stock": stock_stats, "sales": sales_stats,
        "output_rows": len(result), "output_products": int(result.product_code.nunique()),
        "codes_only_in_stock": sorted(set(stock_products.index) - set(sales_products.index)),
        "codes_only_in_sales": sorted(set(sales_products.index) - set(stock_products.index)),
        "codes_without_unit": products.index[products.unit.fillna("").eq("")].tolist(),
        "issue_counts": dict(Counter(issue["issue"] for issue in issues)),
        "notes": [
            "Пустые и некорректные количества сохранены как пропуски, реальные нули сохранены.",
            "Отрицательные значения сохранены, а не превращены в продажи с положительным знаком.",
            "Остатки относятся к началу месяца. Это не текущие остатки и не дни отсутствия товара.",
            "Полнота последнего месяца и склад не определяются из этих файлов. Уточните их отдельно.",
            "Единицы измерения продаж нужно сверить с единицами из файла остатков.",
        ],
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    result.to_csv(output_dir / "monthly_clean.csv", index=False, encoding="utf-8-sig", na_rep="")
    pd.DataFrame(issues, columns=ISSUE_COLUMNS).to_csv(
        output_dir / "quality_issues.csv", index=False, encoding="utf-8-sig"
    )
    (output_dir / "quality_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stock", type=Path, default=ROOT / "IEK" / STOCK_FILE)
    parser.add_argument("--sales", type=Path, default=ROOT / "IEK" / SALES_FILE)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data" / "clean" / "iek")
    parser.add_argument("--sheet", help="Имя листа в обоих файлах; по умолчанию первый лист")
    args = parser.parse_args()
    try:
        report = clean_datasets(args.stock, args.sales, args.output_dir, args.sheet)
    except (ValueError, OSError, KeyError) as error:
        parser.exit(1, f"Ошибка: {error}\n")
    print(f"Готово: {report['output_rows']:,} строк, {report['output_products']:,} товаров.")
    print(f"Результаты: {args.output_dir.resolve()}")
    print(f"Замечания: {report['issue_counts']}")


if __name__ == "__main__":
    main()
